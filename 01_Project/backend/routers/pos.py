from io import BytesIO
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session, selectinload

import models
import schemas
from database import get_db

router = APIRouter(
    prefix="/api/pos",
    tags=["POS"]
)

# Master Excel file for all bills — lives in 02_Datasets
DATASETS_DIR = Path(__file__).resolve().parent.parent.parent.parent / "02_Datasets"
MASTER_EXCEL = DATASETS_DIR / "all_bills.xlsx"

EXCEL_HEADERS = [
    "bill_id", "bill_date", "product_id", "product_name",
    "quantity", "unit_price", "line_total", "bill_total"
]


def _append_bill_to_excel(bill):
    """Append bill rows to the master Excel file, creating it if needed."""
    try:
        DATASETS_DIR.mkdir(parents=True, exist_ok=True)

        if MASTER_EXCEL.exists():
            wb = load_workbook(MASTER_EXCEL)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "Bills"
            ws.append(EXCEL_HEADERS)

        bill_date = str(bill.created_at) if bill.created_at else ""
        for item in bill.items:
            ws.append([
                bill.id,
                bill_date,
                item.product_id,
                item.product_name,
                item.quantity,
                item.unit_price,
                item.line_total,
                bill.total_amount,
            ])

        wb.save(MASTER_EXCEL)
    except Exception as e:
        # Don't fail the bill creation if Excel export fails
        print(f"[WARN] Failed to append bill to master Excel: {e}")


@router.post("/bills", response_model=schemas.BillResponse)
def create_bill(payload: schemas.BillCreate, db: Session = Depends(get_db)):
    if not payload.items:
        raise HTTPException(status_code=400, detail="Bill must contain at least one item")

    product_ids = [item.product_id for item in payload.items]
    products = db.query(models.Product).filter(models.Product.id.in_(product_ids)).all()
    product_map = {product.id: product for product in products}

    bill_items_to_create = []
    total_amount = 0.0

    for item in payload.items:
        if item.quantity <= 0:
            raise HTTPException(status_code=400, detail="Quantity must be greater than zero")

        product = product_map.get(item.product_id)
        if product is None:
            raise HTTPException(status_code=404, detail=f"Product {item.product_id} not found")

        if product.quantity_in_stock < item.quantity:
            raise HTTPException(
                status_code=400,
                detail=f"Insufficient stock for '{product.name}'. Available: {product.quantity_in_stock}"
            )

        line_total = product.price * item.quantity
        total_amount += line_total
        bill_items_to_create.append((product, item.quantity, line_total))

    bill = models.Bill(total_amount=total_amount)
    db.add(bill)
    db.flush()

    for product, quantity, line_total in bill_items_to_create:
        product.quantity_in_stock -= quantity
        db_item = models.BillItem(
            bill_id=bill.id,
            product_id=product.id,
            product_name=product.name,
            quantity=quantity,
            unit_price=product.price,
            line_total=line_total,
        )
        db.add(db_item)

    db.commit()

    created_bill = (
        db.query(models.Bill)
        .options(selectinload(models.Bill.items))
        .filter(models.Bill.id == bill.id)
        .first()
    )

    # Auto-export to master Excel
    _append_bill_to_excel(created_bill)

    return created_bill


@router.get("/bills/recent", response_model=list[schemas.BillSummaryResponse])
def get_recent_bills(limit: int = Query(default=20, ge=1, le=100), db: Session = Depends(get_db)):
    return db.query(models.Bill).order_by(models.Bill.id.desc()).limit(limit).all()


@router.get("/bills/{bill_id}", response_model=schemas.BillResponse)
def get_bill(bill_id: int, db: Session = Depends(get_db)):
    bill = (
        db.query(models.Bill)
        .options(selectinload(models.Bill.items))
        .filter(models.Bill.id == bill_id)
        .first()
    )
    if bill is None:
        raise HTTPException(status_code=404, detail="Bill not found")
    return bill


@router.get("/bills/{bill_id}/excel")
def export_bill_excel(bill_id: int, db: Session = Depends(get_db)):
    bill = (
        db.query(models.Bill)
        .options(selectinload(models.Bill.items))
        .filter(models.Bill.id == bill_id)
        .first()
    )
    if bill is None:
        raise HTTPException(status_code=404, detail="Bill not found")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bill"

    sheet.append(["Bill ID", bill.id])
    sheet.append(["Created At", str(bill.created_at)])
    sheet.append(["Total Amount", round(bill.total_amount, 2)])
    sheet.append([])
    sheet.append(["Item ID", "Product ID", "Product Name", "Quantity", "Unit Price", "Line Total"])

    for item in bill.items:
        sheet.append([
            item.id,
            item.product_id,
            item.product_name,
            item.quantity,
            item.unit_price,
            item.line_total,
        ])

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)

    filename = f"bill_{bill.id}.xlsx"
    headers = {"Content-Disposition": f"attachment; filename={filename}"}

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )

